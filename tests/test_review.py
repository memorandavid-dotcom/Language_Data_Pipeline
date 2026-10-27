import pytest
from openpyxl import load_workbook

from wika_pipeline.review import (
    COLUMNS,
    DECISIONS,
    entry_to_row,
    read_review_workbook,
    review_path,
    select_entries,
    write_review_workbook,
)


def entry(word="abay", meaning="side", flags=None, page=18, **extra):
    base = {
        "entry_id": f"id-{word}",
        "word": word,
        "meaning": meaning,
        "source_page": page,
        "review_flags": flags if flags is not None else [],
        "original_text": f"{word}\n{meaning}",
        "parent_headword": None,
        "homograph_number": None,
        "pronunciation_guide": None,
        "validation_review_status": "extracted",
    }
    base.update(extra)
    return base


def test_only_flagged_entries_are_selected_by_default():
    entries = [entry("a", flags=["no_gloss"]), entry("b"), entry("c", flags=["complex_headword"])]
    assert [e["word"] for e in select_entries(entries)] == ["a", "c"]
    assert [e["word"] for e in select_entries(entries, include_all=True)] == ["a", "b", "c"]


def test_row_shows_flags_original_text_and_status():
    row = entry_to_row(entry("abay", "side", flags=["no_gloss", "cross_reference"]))
    assert row["review_flags"] == "no_gloss; cross_reference"
    assert row["original_text"] == "abay | side"
    assert row["status"] == "extracted"
    assert row["decision"] is None and row["reviewer_id"] is None


def test_text_that_looks_like_a_formula_is_kept_as_text(tmp_path):
    rows = [
        entry_to_row(entry("-n", "= aeropláno, q.v.")),
        entry_to_row(entry("+x", "@mention")),
        entry_to_row(entry("añghÉl", "=1+1")),
    ]
    path = tmp_path / "r.xlsx"
    write_review_workbook(rows, path)
    sheet = load_workbook(path).active
    meaning_column = COLUMNS.index("meaning") + 1
    for row_number in (2, 3, 4):
        assert sheet.cell(row=row_number, column=meaning_column).data_type == "s"
    back = read_review_workbook(path)
    assert [r["word"] for r in back] == ["-n", "+x", "añghÉl"]
    assert [r["meaning"] for r in back] == ["= aeropláno, q.v.", "@mention", "=1+1"]


def test_round_trip_keeps_diacritics_and_ids(tmp_path):
    rows = [entry_to_row(entry("ágewágew, inágewágew", "daily, every day", page=20))]
    path = tmp_path / "r.xlsx"
    write_review_workbook(rows, path)
    back = read_review_workbook(path)
    assert back[0]["word"] == "ágewágew, inágewágew"
    assert back[0]["entry_id"] == "id-ágewágew, inágewágew"
    assert back[0]["source_page"] == "20"


def test_reviewer_input_columns_survive_a_reload(tmp_path):
    rows = [entry_to_row(entry("abay"))]
    path = tmp_path / "r.xlsx"
    write_review_workbook(rows, path)
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet.cell(row=2, column=COLUMNS.index("decision") + 1, value="approve")
    sheet.cell(row=2, column=COLUMNS.index("reviewer_id") + 1, value="rev_001")
    workbook.save(path)
    back = read_review_workbook(path)
    assert back[0]["decision"] == "approve"
    assert back[0]["reviewer_id"] == "rev_001"
    assert back[0]["corrected_word"] is None


def test_decision_column_has_a_drop_down_list(tmp_path):
    path = tmp_path / "r.xlsx"
    write_review_workbook([entry_to_row(entry("abay"))], path)
    validations = load_workbook(path).active.data_validations.dataValidation
    assert len(validations) == 1
    assert validations[0].type == "list"
    for decision in DECISIONS:
        assert decision in validations[0].formula1


def test_an_empty_queue_still_makes_a_valid_workbook(tmp_path):
    path = tmp_path / "r.xlsx"
    write_review_workbook([], path)
    assert read_review_workbook(path) == []


def test_missing_columns_are_reported(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.append(["entry_id", "word"])
    path = tmp_path / "bad.xlsx"
    workbook.save(path)
    with pytest.raises(ValueError, match="missing columns"):
        read_review_workbook(path)


def test_no_temp_file_is_left_behind(tmp_path):
    path = tmp_path / "sub" / "r.xlsx"
    write_review_workbook([entry_to_row(entry("abay"))], path)
    assert sorted(p.name for p in path.parent.iterdir()) == ["r.xlsx"]


def test_review_path_is_per_language_and_document(tmp_path):
    assert review_path(tmp_path, "pag", "doc_1").parts[-2:] == ("pag", "doc_1.review.xlsx")