import json

import pytest
from openpyxl import load_workbook

from wika_pipeline.review import COLUMNS
from wika_pipeline.review_apply import (
    ReviewError,
    apply_actions,
    load_reviewers,
    run_apply,
    validate_rows,
)
from wika_pipeline.review_cli import export_queue
from wika_pipeline.schema import DictionaryEntry

NOW = "2026-02-01T00:00:00+00:00"
REVIEWERS = {
    "rev_pag": {"languages": {"pag"}, "active": True},
    "rev_pam": {"languages": {"pam"}, "active": True},
    "rev_off": {"languages": {"pag"}, "active": False},
}


def make_entry(word, meaning="m", status="extracted", flags=None, language="pag", document="doc_a", **extra):
    entry = DictionaryEntry(
        entry_id=f"id-{word}",
        word=word,
        dialect_code=language,
        dialect_name="Pangasinan",
        meaning=meaning,
        original_text=f"{word}\n{meaning}",
        source_document_id=document,
        source_page=18,
        review_flags=flags if flags is not None else ["no_gloss"],
        validation_review_status=status,
        created_at="2026-01-01T00:00:00+00:00",
        **extra,
    )
    return entry.model_dump(mode="json")


def make_row(entry, decision, reviewer="rev_pag", row_number=2, **extra):
    row = {name: None for name in COLUMNS}
    row.update(
        {
            "entry_id": entry["entry_id"],
            "word": entry["word"],
            "meaning": entry["meaning"],
            "decision": decision,
            "reviewer_id": reviewer,
            "_row": row_number,
        }
    )
    row.update(extra)
    return row


def by_id(entries):
    return {entry["entry_id"]: entry for entry in entries}


def validate(entries, rows, reviewers=REVIEWERS):
    return validate_rows(rows, by_id(entries), reviewers)


def test_rows_without_a_decision_are_ignored():
    entry = make_entry("a")
    actions, errors, skipped = validate([entry], [make_row(entry, None)])
    assert (actions, errors, skipped) == ([], [], 0)


def test_valid_decisions_become_actions():
    a, b, c, d = (make_entry(w) for w in "abcd")
    rows = [
        make_row(a, "approve", row_number=2),
        make_row(b, "correct", row_number=3, corrected_meaning="better"),
        make_row(c, "reject", row_number=4, notes="not a word"),
        make_row(d, "unsure", row_number=5, notes="ask elder"),
    ]
    actions, errors, _ = validate([a, b, c, d], rows)
    assert errors == []
    assert [x["decision"] for x in actions] == ["approve", "correct", "reject", "unsure"]


@pytest.mark.parametrize(
    "build, message",
    [
        (lambda e: make_row(e, "approve", reviewer=None), "reviewer_id is empty"),
        (lambda e: make_row(e, "approve", reviewer="nobody"), "not in the reviewers file"),
        (lambda e: make_row(e, "approve", reviewer="rev_off"), "not active"),
        (lambda e: make_row(e, "approve", reviewer="rev_pam"), "not authorised for language 'pag'"),
        (lambda e: make_row(e, "maybe"), "is not one of"),
        (lambda e: make_row(e, "correct"), "needs a corrected_word or corrected_meaning"),
        (lambda e: make_row(e, "correct", corrected_word="a"), "needs a corrected_word or corrected_meaning"),
        (lambda e: make_row(e, "approve", corrected_word="x"), "only allowed with the decision 'correct'"),
        (lambda e: make_row(e, "reject"), "needs a reason in notes"),
        (lambda e: make_row(e, "unsure"), "needs a reason in notes"),
        (lambda e: dict(make_row(e, "approve"), entry_id="missing"), "entry_id not found"),
        (lambda e: dict(make_row(e, "approve"), word="tampered"), "no longer matches"),
        (lambda e: dict(make_row(e, "approve"), meaning="tampered"), "no longer matches"),
    ],
)
def test_invalid_rows_are_reported_with_their_row_number(build, message):
    entry = make_entry("a")
    actions, errors, _ = validate([entry], [build(entry)])
    assert actions == []
    assert len(errors) == 1
    assert errors[0].startswith("row 2:")
    assert message in errors[0]


def test_decision_is_case_insensitive():
    entry = make_entry("a")
    actions, errors, _ = validate([entry], [make_row(entry, "APPROVE")])
    assert errors == [] and actions[0]["decision"] == "approve"


def test_same_entry_twice_is_an_error():
    entry = make_entry("a")
    _, errors, _ = validate([entry], [make_row(entry, "approve", row_number=2), make_row(entry, "approve", row_number=3)])
    assert any("appears twice" in message for message in errors)


def test_a_final_status_cannot_be_changed_by_a_different_decision():
    entry = make_entry("a", status="approved", reviewer_id="rev_pag")
    _, errors, _ = validate([entry], [make_row(entry, "reject", notes="changed my mind")])
    assert any("already has the final status 'approved'" in message for message in errors)


def test_the_same_decision_again_is_skipped_not_repeated():
    entry = make_entry("a", status="approved", reviewer_id="rev_pag")
    actions, errors, skipped = validate([entry], [make_row(entry, "approve")])
    assert (actions, errors, skipped) == ([], [], 1)


def test_the_same_decision_by_another_reviewer_is_not_skipped():
    entry = make_entry("a", status="approved", reviewer_id="rev_pag")
    other = dict(REVIEWERS, rev_two={"languages": {"pag"}, "active": True})
    _, errors, _ = validate([entry], [make_row(entry, "approve", reviewer="rev_two")], other)
    assert errors


def test_a_correction_already_in_the_entry_is_recognised_as_applied():
    entry = make_entry("fixed", status="approved", reviewer_id="rev_pag")
    row = make_row(entry, "correct", word="original", corrected_word="fixed")
    actions, errors, skipped = validate([entry], [row])
    assert (actions, errors, skipped) == ([], [], 1)


def test_needs_correction_entries_can_be_decided_again():
    entry = make_entry("a", status="needs_correction", reviewer_id="rev_pam")
    actions, errors, _ = validate([entry], [make_row(entry, "approve")])
    assert errors == [] and len(actions) == 1


def test_apply_sets_status_reviewer_time_and_notes():
    entry = make_entry("a")
    actions, _, _ = validate([entry], [make_row(entry, "approve", notes="checked with elder")])
    updated, events = apply_actions([entry], actions, NOW, {"batch_id": "b1", "workbook": "w.xlsx", "workbook_sha256": "h"})
    result = updated[0]
    assert result["validation_review_status"] == "approved"
    assert result["reviewer_id"] == "rev_pag"
    assert result["reviewed_at"] == NOW
    assert result["review_notes"] == "checked with elder"
    assert entry["validation_review_status"] == "extracted"
    assert events[0]["from_status"] == "extracted" and events[0]["to_status"] == "approved"
    assert events[0]["batch_id"] == "b1" and events[0]["workbook_sha256"] == "h"


def test_correct_changes_text_but_keeps_id_and_original_text():
    entry = make_entry("agewagew,", "daily")
    row = make_row(entry, "correct", corrected_word="agewagew", corrected_meaning="every day")
    actions, _, _ = validate([entry], [row])
    updated, events = apply_actions([entry], actions, NOW, {"batch_id": "b", "workbook": "w", "workbook_sha256": "h"})
    result = updated[0]
    assert result["word"] == "agewagew" and result["meaning"] == "every day"
    assert result["entry_id"] == "id-agewagew,"
    assert result["original_text"] == "agewagew,\ndaily"
    assert result["validation_review_status"] == "approved"
    assert events[0]["changes"] == {"word": ["agewagew,", "agewagew"], "meaning": ["daily", "every day"]}


def test_reject_and_unsure_map_to_their_statuses():
    a, b = make_entry("a"), make_entry("b")
    rows = [make_row(a, "reject", notes="x", row_number=2), make_row(b, "unsure", notes="y", row_number=3)]
    actions, _, _ = validate([a, b], rows)
    updated, _ = apply_actions([a, b], actions, NOW, {"batch_id": "b", "workbook": "w", "workbook_sha256": "h"})
    assert [e["validation_review_status"] for e in updated] == ["rejected", "needs_correction"]


def test_load_reviewers_reads_languages_and_active_flag(tmp_path):
    path = tmp_path / "reviewers.yml"
    path.write_text("reviewers:\n  r1:\n    languages: [pag, pam]\n  r2:\n    languages: [ilo]\n    active: false\n", encoding="utf-8")
    reviewers = load_reviewers(path)
    assert reviewers["r1"] == {"languages": {"pag", "pam"}, "active": True}
    assert reviewers["r2"]["active"] is False


@pytest.mark.parametrize(
    "text",
    ["", "reviewers: {}\n", "reviewers:\n  r1: {}\n", "reviewers:\n  r1:\n    languages: []\n", "other: 1\n"],
)
def test_bad_reviewer_files_are_refused(tmp_path, text):
    path = tmp_path / "reviewers.yml"
    path.write_text(text, encoding="utf-8")
    with pytest.raises(ValueError):
        load_reviewers(path)


@pytest.fixture
def workspace(tmp_path):
    entries = [make_entry("alpha", "one"), make_entry("beta", "two"), make_entry("gamma", "three")]
    staging = tmp_path / "staging" / "doc_a.entries.jsonl"
    staging.parent.mkdir()
    with open(staging, "w", encoding="utf-8") as f:
        for entry in entries:
            f.write(json.dumps(entry, ensure_ascii=False) + "\n")
    reviewers = tmp_path / "reviewers.yml"
    reviewers.write_text("reviewers:\n  rev_pag:\n    languages: [pag]\n", encoding="utf-8")
    review_dir = tmp_path / "review"
    workbook, _, _, _ = export_queue(staging, review_dir)
    return {"staging": staging, "reviewers": reviewers, "review_dir": review_dir, "workbook": workbook, "entries": entries}


def fill(workbook_path, decisions):
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    header = {cell.value: cell.column for cell in sheet[1]}
    for row_number, values in decisions.items():
        for column, value in values.items():
            sheet.cell(row=row_number, column=header[column], value=value)
    workbook.save(workbook_path)


def read_jsonl(path):
    with open(path, encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def apply(workspace, **kwargs):
    return run_apply(
        workspace["workbook"], workspace["staging"], workspace["reviewers"], workspace["review_dir"], now=NOW, **kwargs
    )


def test_end_to_end_apply_writes_reviewed_file_and_audit_log(workspace):
    fill(
        workspace["workbook"],
        {
            2: {"decision": "approve", "reviewer_id": "rev_pag"},
            3: {"decision": "correct", "reviewer_id": "rev_pag", "corrected_meaning": "2"},
        },
    )
    summary = apply(workspace)
    assert summary["to_apply"] == 2 and summary["rows_without_decision"] == 1
    reviewed = read_jsonl(workspace["workbook"].parent / "doc_a.reviewed.jsonl")
    assert [e["validation_review_status"] for e in reviewed] == ["approved", "approved", "extracted"]
    assert reviewed[1]["meaning"] == "2"
    audit = read_jsonl(workspace["review_dir"] / "audit_log.jsonl")
    assert [event["decision"] for event in audit] == ["approve", "correct"]
    assert len({event["batch_id"] for event in audit}) == 1
    assert audit[0]["workbook_sha256"] and len(audit[0]["workbook_sha256"]) == 64
    staged = read_jsonl(workspace["staging"])
    assert all(e["validation_review_status"] == "extracted" for e in staged)
    leftovers = [p.name for p in workspace["workbook"].parent.iterdir() if p.name.endswith(".tmp")]
    assert leftovers == []


def test_one_bad_row_blocks_the_whole_batch(workspace):
    fill(
        workspace["workbook"],
        {
            2: {"decision": "approve", "reviewer_id": "rev_pag"},
            3: {"decision": "approve", "reviewer_id": "stranger"},
        },
    )
    with pytest.raises(ReviewError) as caught:
        apply(workspace)
    assert any("row 3" in message for message in caught.value.errors)
    assert not (workspace["workbook"].parent / "doc_a.reviewed.jsonl").exists()
    assert not (workspace["review_dir"] / "audit_log.jsonl").exists()


def test_dry_run_changes_nothing(workspace):
    fill(workspace["workbook"], {2: {"decision": "approve", "reviewer_id": "rev_pag"}})
    summary = apply(workspace, dry_run=True)
    assert summary["dry_run"] is True and summary["to_apply"] == 1
    assert not (workspace["workbook"].parent / "doc_a.reviewed.jsonl").exists()
    assert not (workspace["review_dir"] / "audit_log.jsonl").exists()


def test_applying_the_same_workbook_twice_does_not_duplicate_anything(workspace):
    fill(workspace["workbook"], {2: {"decision": "approve", "reviewer_id": "rev_pag"}})
    apply(workspace)
    second = apply(workspace)
    assert second["to_apply"] == 0 and second["already_applied"] == 1
    assert len(read_jsonl(workspace["review_dir"] / "audit_log.jsonl")) == 1


def test_more_decisions_can_be_added_later_in_the_same_workbook(workspace):
    fill(workspace["workbook"], {2: {"decision": "approve", "reviewer_id": "rev_pag"}})
    apply(workspace)
    fill(workspace["workbook"], {3: {"decision": "reject", "reviewer_id": "rev_pag", "notes": "not a word"}})
    second = apply(workspace)
    assert second["to_apply"] == 1 and second["already_applied"] == 1
    reviewed = read_jsonl(workspace["workbook"].parent / "doc_a.reviewed.jsonl")
    assert [e["validation_review_status"] for e in reviewed] == ["approved", "rejected", "extracted"]
    assert len(read_jsonl(workspace["review_dir"] / "audit_log.jsonl")) == 2


def test_workbook_for_another_document_is_refused(workspace, tmp_path):
    other = tmp_path / "review" / "pag" / "doc_b.review.xlsx"
    other.write_bytes(workspace["workbook"].read_bytes())
    with pytest.raises(ValueError, match="belong to"):
        run_apply(other, workspace["staging"], workspace["reviewers"], workspace["review_dir"], now=NOW)


def test_wrongly_named_workbook_is_refused(workspace, tmp_path):
    bad = tmp_path / "notes.xlsx"
    bad.write_bytes(workspace["workbook"].read_bytes())
    with pytest.raises(ValueError, match="expected a file named"):
        run_apply(bad, workspace["staging"], workspace["reviewers"], workspace["review_dir"], now=NOW)