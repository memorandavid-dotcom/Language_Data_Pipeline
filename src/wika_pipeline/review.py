from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

DECISIONS = ("approve", "correct", "reject", "unsure")

READ_ONLY_COLUMNS = [
    "entry_id",
    "source_page",
    "word",
    "parent_headword",
    "homograph_number",
    "pronunciation_guide",
    "meaning",
    "review_flags",
    "original_text",
    "status",
]
INPUT_COLUMNS = [
    "decision",
    "corrected_word",
    "corrected_meaning",
    "reviewer_id",
    "notes",
]
COLUMNS = READ_ONLY_COLUMNS + INPUT_COLUMNS
SHEET_NAME = "review"
WIDTHS = {
    "entry_id": 38,
    "source_page": 8,
    "word": 22,
    "parent_headword": 18,
    "homograph_number": 8,
    "pronunciation_guide": 18,
    "meaning": 50,
    "review_flags": 28,
    "original_text": 50,
    "status": 14,
    "decision": 12,
    "corrected_word": 22,
    "corrected_meaning": 50,
    "reviewer_id": 14,
    "notes": 30,
}


def entry_to_row(entry):
    lines = (entry.get("original_text") or "").split("\n")
    return {
        "entry_id": entry["entry_id"],
        "source_page": entry.get("source_page"),
        "word": entry["word"],
        "parent_headword": entry.get("parent_headword"),
        "homograph_number": entry.get("homograph_number"),
        "pronunciation_guide": entry.get("pronunciation_guide"),
        "meaning": entry.get("meaning"),
        "review_flags": "; ".join(entry.get("review_flags") or []),
        "original_text": " | ".join(line for line in lines if line),
        "status": entry.get("validation_review_status"),
        "decision": None,
        "corrected_word": None,
        "corrected_meaning": None,
        "reviewer_id": None,
        "notes": None,
    }


def select_entries(entries, include_all=False):
    if include_all:
        return list(entries)
    return [entry for entry in entries if entry.get("review_flags")]


def write_review_workbook(rows, path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    bold = Font(bold=True)
    for column, name in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column, value=name)
        cell.font = bold
        sheet.column_dimensions[cell.column_letter].width = WIDTHS[name]
    for row_number, row in enumerate(rows, start=2):
        for column, name in enumerate(COLUMNS, start=1):
            value = row.get(name)
            if value is None or value == "":
                continue
            cell = sheet.cell(row=row_number, column=column)
            cell.value = str(value)
            cell.data_type = "s"
    sheet.freeze_panes = "D2"
    decision_column = sheet.cell(row=1, column=COLUMNS.index("decision") + 1).column_letter
    validation = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
    validation.add(f"{decision_column}2:{decision_column}{max(len(rows) + 1, 2)}")
    sheet.add_data_validation(validation)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    temp = Path(str(path) + ".tmp")
    workbook.save(temp)
    temp.replace(path)


def read_review_workbook(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            raise ValueError(f"{path}: the workbook is empty")
        names = [str(value).strip() if value is not None else "" for value in header]
        missing = [name for name in COLUMNS if name not in names]
        if missing:
            raise ValueError(f"{path}: missing columns {missing}")
        rows = []
        for row_number, values in enumerate(iterator, start=2):
            row = {}
            for name, value in zip(names, values):
                if name in COLUMNS:
                    text = None if value is None else str(value).strip()
                    row[name] = text or None
            if any(row.get(name) for name in COLUMNS):
                row["_row"] = row_number
                rows.append(row)
        return rows
    finally:
        workbook.close()


def review_path(review_dir, language_code, document_id):
    return Path(review_dir) / language_code / f"{document_id}.review.xlsx"