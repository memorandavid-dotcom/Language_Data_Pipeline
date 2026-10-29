import json

import pytest
from openpyxl import load_workbook

from wika_pipeline.export import approval_basis
from wika_pipeline.review import read_review_workbook
from wika_pipeline.review_apply import run_apply
from wika_pipeline.sampling import (
    SamplingError,
    conclude_sample,
    create_full_review,
    create_sample,
    ids_hash,
    load_policy,
    population_of,
    required_sample_size,
    stratified_pick,
)
from wika_pipeline.schema import DictionaryEntry

NOW = "2026-04-01T00:00:00+00:00"
POLICY = {"max_error_rate": 0.05, "confidence": 0.95}
SIZE = 59


def make_entry(index, flagged=False, status="extracted"):
    return DictionaryEntry(
        entry_id=f"id-{index:04d}",
        word=f"word{index:04d}",
        dialect_code="pag",
        dialect_name="Pangasinan",
        meaning=f"meaning {index}",
        original_text=f"word{index:04d}\nmeaning {index}",
        source_document_id="doc_a",
        source_page=index // 10 + 1,
        review_flags=["complex_headword"] if flagged else [],
        validation_review_status=status,
        created_at="2026-01-01T00:00:00+00:00",
    ).model_dump(mode="json")


@pytest.fixture
def project(tmp_path):
    entries = [make_entry(i, flagged=(i % 30 == 0)) for i in range(300)]
    staging = tmp_path / "staging" / "doc_a.entries.jsonl"
    staging.parent.mkdir()
    staging.write_text("".join(json.dumps(e) + "\n" for e in entries), encoding="utf-8")
    reviewers = tmp_path / "reviewers.yml"
    reviewers.write_text(
        "reviewers:\n  rev_pag:\n    languages: [pag]\n  rev_off:\n    languages: [pag]\n    active: false\n  rev_pam:\n    languages: [pam]\n",
        encoding="utf-8",
    )
    return {"staging": staging, "reviewers": reviewers, "review_dir": tmp_path / "review", "entries": entries}


def sample(project, seed=7):
    return create_sample(project["staging"], project["review_dir"], POLICY, seed=seed, now=NOW)


def decide(project, result, overrides=None, reviewer="rev_pag"):
    workbook = result["workbook"]
    book = load_workbook(workbook)
    sheet = book.active
    header = {cell.value: cell.column for cell in sheet[1]}
    overrides = overrides or {}
    for row_number in range(2, sheet.max_row + 1):
        extra = overrides.get(row_number, {"decision": "approve"})
        values = dict(extra, reviewer_id=reviewer)
        for column, value in values.items():
            sheet.cell(row=row_number, column=header[column], value=value)
    book.save(workbook)
    return run_apply(workbook, project["staging"], project["reviewers"], project["review_dir"], now=NOW)


def conclude(project, result, approver="rev_pag", **kwargs):
    return conclude_sample(
        result["manifest_path"], approver, project["reviewers"], project["review_dir"], now=NOW, **kwargs
    )


def read_jsonl(path):
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def reviewed(project):
    return read_jsonl(project["review_dir"] / "pag" / "doc_a.reviewed.jsonl")


def test_sample_sizes_follow_the_zero_error_rule():
    assert required_sample_size(0.02, 0.95) == 149
    assert required_sample_size(0.01, 0.95) == 299
    assert required_sample_size(0.05, 0.95) == 59
    assert required_sample_size(0.05, 0.99) > required_sample_size(0.05, 0.95)


def test_policy_file_rules(tmp_path):
    path = tmp_path / "p.yml"
    path.write_text("max_error_rate: 0.02\nconfidence: 0.95\n", encoding="utf-8")
    assert load_policy(path) == {"max_error_rate": 0.02, "confidence": 0.95}
    for text in ["", "max_error_rate: 0.02\n", "max_error_rate: 0\nconfidence: 0.9\n", "max_error_rate: 0.1\nconfidence: 1\n", "max_error_rate: high\nconfidence: 0.9\n"]:
        path.write_text(text, encoding="utf-8")
        with pytest.raises(ValueError):
            load_policy(path)


def test_population_is_unflagged_and_still_open():
    entries = [make_entry(1), make_entry(2, flagged=True), make_entry(3, status="approved"), make_entry(4, status="needs_correction")]
    assert [e["entry_id"] for e in population_of(entries)] == ["id-0001", "id-0004"]


def test_stratified_pick_takes_one_per_slice_and_is_reproducible():
    ids = [f"i{n}" for n in range(1000)]
    first = stratified_pick(ids, 10, seed=1)
    assert first == stratified_pick(ids, 10, seed=1)
    assert first != stratified_pick(ids, 10, seed=2)
    for index, pick in enumerate(first):
        assert index * 100 <= ids.index(pick) < (index + 1) * 100


def test_sample_covers_only_unflagged_entries_and_writes_a_workbook_and_record(project):
    result = sample(project)
    rows = read_review_workbook(result["workbook"])
    assert len(rows) == SIZE
    flagged = {e["entry_id"] for e in project["entries"] if e["review_flags"]}
    assert not flagged & {row["entry_id"] for row in rows}
    manifest = json.loads(result["manifest_path"].read_text(encoding="utf-8"))
    assert manifest["status"] == "open" and manifest["seed"] == 7
    assert manifest["population_size"] == 290 and manifest["sample_size"] == SIZE
    assert manifest["population_ids_sha256"] == ids_hash(manifest["population_ids"])
    assert set(manifest["sample_ids"]) <= set(manifest["population_ids"])
    assert result["workbook"].name == f"doc_a.audit-{manifest['sample_id']}.review.xlsx"


def test_same_seed_gives_the_same_sample(project, tmp_path):
    first = sample(project)
    other = {**project, "review_dir": tmp_path / "second"}
    second = create_sample(other["staging"], other["review_dir"], POLICY, seed=7, now=NOW)
    assert first["manifest"]["sample_ids"] == second["manifest"]["sample_ids"]
    assert first["manifest"]["sample_id"] == second["manifest"]["sample_id"]


def test_a_second_sample_is_refused_while_one_is_open(project):
    sample(project, seed=1)
    with pytest.raises(ValueError, match="still open"):
        sample(project, seed=2)


def test_a_population_too_small_for_sampling_is_refused(project):
    tiny = [make_entry(i) for i in range(40)]
    project["staging"].write_text("".join(json.dumps(e) + "\n" for e in tiny), encoding="utf-8")
    with pytest.raises(ValueError, match="reviewing all of them"):
        sample(project)


def test_concluding_before_decisions_are_applied_is_refused(project):
    result = sample(project)
    with pytest.raises(ValueError, match="No decisions have been applied"):
        conclude(project, result)


def test_incomplete_sample_lists_the_undecided_entries(project):
    result = sample(project)
    decide(project, result, overrides={row: {} for row in range(3, SIZE + 2)})
    with pytest.raises(SamplingError) as caught:
        conclude(project, result)
    assert len(caught.value.problems) == SIZE - 1
    assert all("no decision applied yet" in p for p in caught.value.problems)


def test_clean_sample_approves_the_rest_in_bulk_with_an_audit_trail(project):
    result = sample(project)
    decide(project, result)
    outcome = conclude(project, result)
    assert outcome["outcome"] == "passed" and outcome["approved"] == 290 - SIZE
    entries = reviewed(project)
    by_status = {}
    for entry in entries:
        by_status.setdefault(entry["validation_review_status"], []).append(entry)
    assert len(by_status["approved"]) == 290
    assert len(by_status["extracted"]) == 10
    assert all(e["review_flags"] for e in by_status["extracted"])
    bulk = [e for e in by_status["approved"] if "audit sample" in (e["review_notes"] or "")]
    assert len(bulk) == 290 - SIZE
    assert all(e["reviewer_id"] == "rev_pag" for e in bulk)
    audit = read_jsonl(project["review_dir"] / "audit_log.jsonl")
    bulk_events = [e for e in audit if e["decision"] == "bulk_approve"]
    assert len(bulk_events) == 290 - SIZE
    assert all(e["audit_sample_id"] == result["manifest"]["sample_id"] for e in bulk_events)
    manifest = json.loads(result["manifest_path"].read_text(encoding="utf-8"))
    assert manifest["status"] == "passed" and manifest["bulk_approved"] == 290 - SIZE


def test_one_error_in_the_sample_fails_it_and_approves_nothing_in_bulk(project):
    result = sample(project)
    decide(project, result, overrides={5: {"decision": "reject", "notes": "not a word"}})
    outcome = conclude(project, result)
    assert outcome["outcome"] == "failed" and outcome["errors_found"] == 1
    statuses = [e["validation_review_status"] for e in reviewed(project)]
    assert statuses.count("approved") == SIZE - 1
    assert json.loads(result["manifest_path"].read_text(encoding="utf-8"))["status"] == "failed"


def test_a_correction_counts_as_an_error(project):
    result = sample(project)
    decide(project, result, overrides={4: {"decision": "correct", "corrected_meaning": "better"}})
    assert conclude(project, result)["outcome"] == "failed"


def test_dry_run_changes_nothing(project):
    result = sample(project)
    decide(project, result)
    before = (project["review_dir"] / "pag" / "doc_a.reviewed.jsonl").read_text(encoding="utf-8")
    outcome = conclude(project, result, dry_run=True)
    assert outcome["outcome"] == "passed" and outcome["to_approve"] == 290 - SIZE
    assert (project["review_dir"] / "pag" / "doc_a.reviewed.jsonl").read_text(encoding="utf-8") == before
    assert json.loads(result["manifest_path"].read_text(encoding="utf-8"))["status"] == "open"


@pytest.mark.parametrize("approver", ["nobody", "rev_off", "rev_pam"])
def test_only_an_active_authorised_reviewer_can_conclude(project, approver):
    result = sample(project)
    decide(project, result)
    with pytest.raises(ValueError, match="active reviewer authorised"):
        conclude(project, result, approver=approver)


def test_a_concluded_sample_cannot_be_concluded_again(project):
    result = sample(project)
    decide(project, result)
    conclude(project, result)
    with pytest.raises(ValueError, match="already passed"):
        conclude(project, result)


def test_an_edited_manifest_is_detected(project):
    result = sample(project)
    decide(project, result)
    manifest = json.loads(result["manifest_path"].read_text(encoding="utf-8"))
    manifest["population_ids"] = manifest["population_ids"][:-5]
    result["manifest_path"].write_text(json.dumps(manifest), encoding="utf-8")
    with pytest.raises(ValueError, match="edited"):
        conclude(project, result)


def test_a_shrunken_sample_is_detected(project):
    result = sample(project)
    decide(project, result)
    manifest = json.loads(result["manifest_path"].read_text(encoding="utf-8"))
    manifest["sample_size"] = 10
    manifest["sample_ids"] = manifest["sample_ids"][:10]
    result["manifest_path"].write_text(json.dumps(manifest), encoding="utf-8")
    with pytest.raises(ValueError, match="smaller than the policy"):
        conclude(project, result)


def test_full_review_after_a_failed_sample_lists_every_remaining_unflagged_entry(project):
    result = sample(project)
    decide(project, result, overrides={5: {"decision": "reject", "notes": "bad"}})
    conclude(project, result)
    workbook, count = create_full_review(result["manifest_path"], project["review_dir"])
    assert count == 290 - SIZE
    assert len(read_review_workbook(workbook)) == count
    assert workbook.name == f"doc_a.full-{result['manifest']['sample_id']}.review.xlsx"
    with pytest.raises(FileExistsError):
        create_full_review(result["manifest_path"], project["review_dir"])


def test_full_review_is_refused_unless_the_sample_failed(project):
    result = sample(project)
    with pytest.raises(ValueError, match="only needed after a failed sample"):
        create_full_review(result["manifest_path"], project["review_dir"])


def test_full_review_workbook_can_be_applied_like_any_other(project):
    result = sample(project)
    decide(project, result, overrides={5: {"decision": "reject", "notes": "bad"}})
    conclude(project, result)
    workbook, _ = create_full_review(result["manifest_path"], project["review_dir"])
    book = load_workbook(workbook)
    sheet = book.active
    header = {cell.value: cell.column for cell in sheet[1]}
    for row_number in range(2, 5):
        sheet.cell(row=row_number, column=header["decision"], value="approve")
        sheet.cell(row=row_number, column=header["reviewer_id"], value="rev_pag")
    book.save(workbook)
    summary = run_apply(workbook, project["staging"], project["reviewers"], project["review_dir"], now=NOW)
    assert summary["to_apply"] == 3


def test_export_manifest_can_tell_bulk_approvals_from_individual_ones():
    approved = [{"entry_id": "a"}, {"entry_id": "b"}, {"entry_id": "c"}]
    events = [
        {"entry_id": "a", "to_status": "approved", "decision": "approve"},
        {"entry_id": "b", "to_status": "approved", "decision": "bulk_approve"},
        {"entry_id": "c", "to_status": "approved", "decision": "bulk_approve"},
    ]
    assert approval_basis(approved, events) == {"approve": 1, "bulk_approve": 2}
    assert approval_basis(approved, []) == {"unknown": 3}
